import docx,json,re
from docx.text.paragraph import Paragraph
from docx.table import Table
d=docx.Document('c47/in.docx')
els=[]; started=False; sec=None
def segs(p):
    o=[[r.text,bool(r.bold),bool(r.italic)] for r in p.runs if r.text!='']
    return o or [[p.text,False,False]]
P=lambda k,s: dict(type='para',kind=k,segs=s,p=0,x0=0)
def strip(s,pat):
    s=[x[:] for x in s]; s[0][0]=re.sub(pat,'',s[0][0],count=1); return s
for ch in d.element.body.iterchildren():
    tag=ch.tag.split('}')[1]
    if tag=='p':
        p=Paragraph(ch,d); t=p.text.strip(); st=p.style.name
        if not t: continue
        if not started:
            if st.startswith('Heading') and 'What You Will Learn' in t: started=True
            else: continue
        s=segs(p)
        if st.startswith('Heading'):
            m=re.match(r'^(\d+)\.\s+(.*)',t); n=int(m.group(1)); title=m.group(2)
            if title=='Study Goal': continue
            if 5<=n<=17: els.append(P('h1',[[f'{n-4}. {title}',False,False]])); sec='T'
            else: els.append(P('h1',[[title,False,False]])); sec=title
            continue
        if st=='List Bullet': els.append(P('bul',s)); continue
        if st=='List Number': els.append(P('num',s)); continue
        if sec=='Scripture-Tracing Exercise' and re.match(r'^\d+\.\s',t): els.append(P('num',strip(s,r'^\d+\.\s*'))); continue
        if sec=='Review and Discussion Questions' and re.match(r'^\d+\.\s',t): els.append(P('num',strip(s,r'^\d+\.\s*'))); continue
        if sec=='Teaching Outline' and re.match(r'^[IVX]+\.\s',t): els.append(P('num',strip(s,r'^[IVX]+\.\s*'))); continue
        els.append(P('body',s))
    elif tag=='tbl' and started:
        tb=Table(ch,d); ps=[q for q in tb.cell(0,0).paragraphs if q.text.strip()]
        els.append(P('ctitle',[[ps[0].text.strip(),True,False]]))
        for q in ps[1:]: els.append(P('cbody',segs(q)))
for e in els:
    for s in e['segs']: s[0]=s[0].replace("'","’")
json.dump(els,open('els47.json','w'),ensure_ascii=False)
for i,e in enumerate(els): print(i,e['kind'],''.join(s[0] for s in e['segs'])[:60])

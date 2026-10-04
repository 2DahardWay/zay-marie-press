import docx,json,re
from docx.text.paragraph import Paragraph
from docx.table import Table
d=docx.Document('c51/fin.docx')
els=[]; started=False; sec=None; tn=0
def segs(p):
    o=[[r.text,bool(r.bold),bool(r.italic)] for r in p.runs if r.text!='']
    return o or [[p.text,False,False]]
P=lambda k,s: dict(type='para',kind=k,segs=s,p=0,x0=0)
def strip(s,pat):
    s=[x[:] for x in s]; s[0][0]=re.sub(pat,'',s[0][0],count=1); return s
for ch in d.element.body.iterchildren():
    tag=ch.tag.split('}')[1]
    if tag=='p':
        p=Paragraph(ch,d); t=p.text.strip(); st=p.style.name
        if not t: continue
        if not started:
            if st.startswith('Heading') and t=='What You Will Learn in This Study': started=True
            else: continue
        if t.startswith('ZAY-MARIE PRESS DIGITAL STUDIES') or t.startswith('©') or t.startswith('Copyright'): continue
        s=segs(p)
        if st=='Heading 1':
            if t=='Study Goal': continue
            sec=t; els.append(P('h1',[[t,False,False]])); continue
        if st=='Heading 2': els.append(P('body',[[t,False,False]])); continue
        if st=='List Bullet':
            if sec=='Teaching Outline': els.append(P('bul',strip(s,r'^[A-Z]\.\s*')))
            else: els.append(P('bul',s))
            continue
        if sec=='Teaching Outline' and re.match(r'^[IVX]+\.\s',t):
            tn+=1; x=strip(s,r'^[IVX]+\.\s*'); x[0][0]=f'{tn}. '+x[0][0]; els.append(P('body',x)); continue
        if re.match(r'^\d+\.\t',p.text) and not st.startswith('Heading'): els.append(P('num',strip(s,r'^\d+\.\s*'))); continue
        els.append(P('body',s))
    elif tag=='tbl' and started:
        tb=Table(ch,d); ps=[q for q in tb.cell(0,0).paragraphs if q.text.strip()]
        els.append(P('ctitle',[[ps[0].text.strip(),True,False]]))
        for q in ps[1:]: els.append(P('cbody',segs(q)))
for e in els:
    for s in e['segs']: s[0]=s[0].replace("'","’")
json.dump(els,open('els51.json','w'),ensure_ascii=False)
for i,e in enumerate(els): print(i,e['kind'],''.join(s[0] for s in e['segs'])[:60])

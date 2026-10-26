import docx,copy
from docx.shared import Pt,Twips,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
NAVY=RGBColor(0x1B,0x33,0x50)
d=docx.Document()
s=d.sections[0]
s.page_width=Twips(9677);s.page_height=Twips(14515)
s.top_margin=Twips(864);s.bottom_margin=Twips(1224);s.left_margin=Twips(1031);s.right_margin=Twips(1031)
s.header_distance=Twips(400);s.footer_distance=Twips(475)
st=d.styles['Normal'];st.font.name='Roboto';st.font.size=Pt(10)
rpr=st.element.get_or_add_rPr();rf=rpr.find(qn('w:rFonts'))
for a in ('w:ascii','w:hAnsi','w:eastAsia','w:cs'): rf.set(qn(a),'Roboto')
def run(p,t,b=False,i=False,size=None,color=None):
    r=p.add_run(t);r.bold=b;r.italic=i
    r.font.name='Roboto';r._r.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Roboto')
    if size:r.font.size=Pt(size)
    r.font.color.rgb=color or RGBColor(0,0,0);return r
def para(after=6,before=0,keep=False):
    p=d.add_paragraph();pf=p.paragraph_format;pf.space_after=Pt(after);pf.space_before=Pt(before);pf.line_spacing=1.08;pf.alignment=WD_ALIGN_PARAGRAPH.LEFT;pf.keep_with_next=keep;return p
p=para(after=8,before=0);run(p,'How to Use the Digital Studies',b=True,size=17,color=NAVY)
p=para();run(p,'These studies are designed to help you examine Scripture carefully, trace its arguments, and distinguish what a passage explicitly states from conclusions that may be drawn from it. They are not intended merely to be read from beginning to end and then set aside.')
items=[('Read the study first.','Begin with the central question and work through the study in order. Follow the Scripture references as the argument develops rather than reading individual passages in isolation.'),
('Keep your Bible open.','Read the cited passages in their immediate context. Pay particular attention to audience, setting, terminology, and what the text explicitly says or does not say.'),
('Observe before concluding.','Do not begin with the conclusion you expect to find. Let the passages establish their own categories, identities, relationships, and distinctions.'),
('Use the review sections.','The Key Distinctions, Scripture-Tracing Exercise, and Review and Discussion Questions are designed to help you test whether you can reproduce the study’s reasoning from Scripture rather than merely remember its conclusion.'),
('Go deeper with the Guided Further Study.','After completing the main study, work through the investigations slowly. Record your observations before consulting the study again. The goal is to develop the ability to examine the biblical evidence independently.'),
('Teach what you have learned.','Use the Teaching Outline to explain the subject in your own words. If you can trace the relevant passages, explain their context, and show how the conclusions arise from the text, you have moved beyond reading the study to understanding it.')]
for a,b in items:
    p=para();run(p,a+' ',b=True);run(p,b)
# closing callout
t=d.add_table(rows=1,cols=1)
tbl=t._tbl;pr=tbl.tblPr
def el(n,**k):
    e=OxmlElement(n)
    for a,v in k.items(): e.set(qn('w:'+a),v)
    return e
pr.append(el('w:tblW',w='7615',type='dxa'));pr.append(el('w:tblInd',w='100',type='dxa'));pr.append(el('w:tblLayout',type='fixed'))
bd=OxmlElement('w:tblBorders')
for side in ('top','left','bottom','right'): bd.append(el('w:'+side,val='single',sz='8',space='0',color='1B3350'))
pr.append(bd)
cm=OxmlElement('w:tblCellMar')
for side in ('top','left','bottom','right'): cm.append(el('w:'+side,w='100',type='dxa'))
pr.append(cm)
g=tbl.find(qn('w:tblGrid'));g[0].set(qn('w:w'),'7615')
c=t.cell(0,0);tcPr=c._tc.get_or_add_tcPr();tcPr.append(el('w:tcW',w='7615',type='dxa'));tcPr.append(el('w:shd',val='clear',color='auto',fill='E3EDF7'))
cp=c.paragraphs[0];cp.paragraph_format.space_after=Pt(2);run(cp,'THE GOAL',b=True,size=10,color=NAVY)
cp2=c.add_paragraph();cp2.paragraph_format.space_after=Pt(0);cp2.paragraph_format.line_spacing=1.08
run(cp2,'The goal is not simply to know the conclusion. It is to understand how Scripture establishes it.',b=True)
# spacing before table: previous paragraph after 10pt
d.paragraphs[-1].paragraph_format.space_after=Pt(10)
ip=para(after=0,before=14)
run(ip,'Copyright © 2026 Alexander S. Marien. All rights reserved. Published by Zay-Marie Press. Free for the personal use of the reader; no part may be reproduced or distributed without permission.',size=8)
# footer
f=s.footer;fp=f.paragraphs[0];fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
run(fp,'ZAY-MARIE PRESS DIGITAL STUDIES SERIES • 1',size=7,color=RGBColor(0x55,0x55,0x55))
d.save('How_to_Use_the_Digital_Studies.docx')

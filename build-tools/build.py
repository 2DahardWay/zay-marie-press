import json,re,copy,sys,importlib
from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn,nsdecls
from docx.oxml import parse_xml
from PIL import Image
cfg=json.load(open(sys.argv[1]))      # els, edits module, cover, title, out, imgs
els=json.load(open(cfg['els']))
ed=importlib.import_module(cfg['edits']); els=ed.apply(els)
d=Document('/tmp/claude-0/-home-claude/c0ed2e22-30fa-5e68-8953-7079a3cff252/scratchpad/tpl/conf.docx'); body=d.element.body; ch=list(body)
proto={k:copy.deepcopy(ch[i]) for k,i in dict(tochead=1,toc=2,h1pb=24,body=25,bul=27,call=35,h2=36,h1=38,sub=43,num=154,imp=179,impc=180).items()}
cover=ch[0]; final=ch[-1]
for e in ch[1:-1]: body.remove(e)
Image.open(cfg['cover']).convert('RGB').save('cover.png')
for rel in d.part.rels.values():
    if 'image' in rel.reltype: rel.target_part._blob=open('cover.png','rb').read()
def clear_runs(pe):
    for r in pe.findall(qn('w:r')): pe.remove(r)
def runrpr(pe):
    r=pe.find(qn('w:r')); return copy.deepcopy(r.find(qn('w:rPr')))
def mkrun(rpr,text,bold=None,ital=None):
    r=parse_xml(f'<w:r {nsdecls("w")}/>'); rp=copy.deepcopy(rpr)
    if bold:
        for e in rp.findall(qn('w:b')): rp.remove(e)
        rp.find(qn('w:rFonts')).addnext(parse_xml(f'<w:b {nsdecls("w")}/>'))
    if ital:
        anchor=rp.find(qn('w:b')) if rp.find(qn('w:b')) is not None else rp.find(qn('w:rFonts'))
        anchor.addnext(parse_xml(f'<w:i {nsdecls("w")}/>'))
    r.append(rp)
    for k,pt in enumerate(text.split('\t')):
        if k>0: r.append(parse_xml(f'<w:tab {nsdecls("w")}/>'))
        t=parse_xml(f'<w:t {nsdecls("w")} xml:space="preserve"/>'); t.text=pt; r.append(t)
    return r
def mkp(key,segs,forcebold=False):
    pe=copy.deepcopy(proto[key]); rpr=runrpr(pe); clear_runs(pe)
    for t,b,i in segs: pe.append(mkrun(rpr,t,bold=(b or forcebold),ital=i))
    return pe
def ptext(e): return ''.join(s[0] for s in e['segs'])
add=lambda e: final.addprevious(e)
def esc(s): return s.replace('&','&amp;').replace('<','&lt;').replace('>','&gt;')
def mktable(rows,xs):
    tot=sum(x1-x0 for x0,x1 in xs); W=[int(round((x1-x0)/tot*7815)) for x0,x1 in xs]; W[-1]=7815-sum(W[:-1])
    bd='<w:tcBorders>'+''.join(f'<w:{s} w:val="single" w:sz="8" w:space="0" w:color="2B4F6B"/>' for s in ('top','left','bottom','right'))+'</w:tcBorders>'
    mar='<w:tcMar><w:top w:w="70" w:type="dxa"/><w:left w:w="100" w:type="dxa"/><w:bottom w:w="70" w:type="dxa"/><w:right w:w="100" w:type="dxa"/></w:tcMar>'
    x=f'<w:tbl {nsdecls("w")}><w:tblPr><w:tblW w:w="7815" w:type="dxa"/><w:tblInd w:w="-100" w:type="dxa"/></w:tblPr><w:tblGrid>'+''.join(f'<w:gridCol w:w="{w}"/>' for w in W)+'</w:tblGrid>'
    for ri,row in enumerate(rows):
        hdr=ri==0
        x+='<w:tr><w:trPr><w:cantSplit/>'+('<w:tblHeader/>' if hdr else '')+'</w:trPr>'
        for ci,cell in enumerate(row):
            fill='3B6E91' if hdr else 'FFFFFF'
            rp='<w:rPr><w:rFonts w:ascii="Roboto" w:hAnsi="Roboto" w:eastAsia="Roboto" w:cs="Roboto"/>'+('<w:b/><w:color w:val="FFFFFF"/>' if hdr else '')+'</w:rPr>'
            x+=f'<w:tc><w:tcPr><w:tcW w:w="{W[ci]}" w:type="dxa"/>{bd}<w:shd w:val="clear" w:color="auto" w:fill="{fill}"/>{mar}</w:tcPr><w:p><w:pPr>'+('<w:keepNext/>' if ri<1 else '')+f'</w:pPr><w:r>{rp}<w:t xml:space="preserve">{esc(cell)}</w:t></w:r></w:p></w:tc>'
        x+='</w:tr>'
    return parse_xml(x+'</w:tbl>')
# TOC
heads=[ptext(e) for e in els if e['type']=='para' and e['kind']=='h1' and ptext(e)!='Study Outline']
add(copy.deepcopy(proto['tochead']))
for h in heads:
    pe=copy.deepcopy(proto['toc']); rpr=runrpr(pe); clear_runs(pe); pe.append(mkrun(rpr,h+'\t0')); add(pe)
first=True; i=0
while i<len(els):
    e=els[i]
    if e['type']=='tbl':
        add(mktable(e['rows'],e['xs'])); i+=1; continue
    if e['type']=='img':
        fp=parse_xml(f'<w:p {nsdecls("w")}><w:pPr><w:keepNext/><w:spacing w:before="120" w:after="120"/><w:jc w:val="center"/></w:pPr></w:p>')
        add(fp); par=Paragraph(fp,d)
        from docx.shared import Pt
        par.add_run().add_picture(cfg['imgs'][str(e['p'])],width=Pt(min(e['w'],340))); i+=1; continue
    k=e['kind']; t=ptext(e)
    if k=='h1' and t=='Study Outline': add(mkp('h2',[[t,True,False]])); i+=1; continue
    if k=='h1':
        add(mkp('h1pb' if first else 'h1',[[t,False,False]])); first=False; i+=1; continue
    if k=='sub': add(mkp('sub',[[t,True,False]])); i+=1; continue
    if k=='bul': add(mkp('bul',e['segs'])); i+=1; continue
    if k=='num': add(mkp('num',e['segs'])); i+=1; continue
    if k=='ctitle':
        tb=copy.deepcopy(proto['call']); tc=tb.find('.//'+qn('w:tc')); ps=tc.findall(qn('w:p'))
        for r in ps[0].findall(qn('w:r')):
            for tt in r.findall(qn('w:t')): tt.text=t
        bproto=ps[1]; tc.remove(ps[1]); j=i+1
        while j<len(els) and els[j]['type']=='para' and els[j]['kind']=='cbody':
            bp=copy.deepcopy(bproto); rpr=copy.deepcopy(bp.find(qn('w:r')).find(qn('w:rPr'))); clear_runs(bp)
            for tx,b,it in els[j]['segs']: bp.append(mkrun(rpr,tx,bold=b,ital=it))
            tc.append(bp); j+=1
        add(tb); i=j; continue
    if k=='ctitle_nobox' or t.startswith('Copyright ©'): i+=1; continue
    if k=='body': add(mkp('body',e['segs'])); i+=1; continue
    i+=1
add(copy.deepcopy(proto['imp'])); add(copy.deepcopy(proto['impc']))
d.core_properties.title=cfg['title']
d.save('out0.docx'); print('saved',len(heads))

import re,copy
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls
d=Document('conf.docx'); n=0
for p in d.paragraphs:
    if re.match(r'^\d\.\d\s',p.text) or p.text.startswith('The table sets out'):
        pPr=p._p.get_or_add_pPr()
        if pPr.find(qn('w:keepNext')) is None: pPr.insert(0,OxmlElement('w:keepNext')); n+=1
body=d.element.body; sp=0
for el in list(body):
    nx=el.getnext()
    if el.tag==qn('w:tbl') and nx is not None and nx.tag==qn('w:tbl'):
        el.addnext(parse_xml(f'<w:p {nsdecls("w")}><w:pPr><w:spacing w:before="0" w:after="0" w:line="200" w:lineRule="exact"/></w:pPr></w:p>')); sp+=1
ps=d.paragraphs
i=[k for k,p in enumerate(ps) if p.text.strip()=='ZAY-MARIE PRESS DIGITAL STUDIES'][0]
for k in (i-2,i-1,i):
    pPr=ps[k]._p.get_or_add_pPr()
    if pPr.find(qn('w:keepNext')) is None: pPr.insert(0,OxmlElement('w:keepNext'))
d.save('conf.docx'); print('keepNext sub',n,'spacers',sp)

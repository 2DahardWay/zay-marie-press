import docx,json,re
from docx.text.paragraph import Paragraph
from docx.table import Table
W='{http://schemas.openxmlformats.org/wordprocessingml/2006/main}w'
d=docx.Document('c50/fin.docx')
els=[]; started=False; sec=None
def segs(p):
    o=[[r.text,bool(r.bold),bool(r.italic)] for r in p.runs if r.text!='']
    return o or [[p.text,False,False]]
P=lambda k,s: dict(type='para',kind=k,segs=s,p=0,x0=0)
def strip(s,pat):
    s=[x[:] for x in s]; s[0][0]=re.sub(pat,'',s[0][0],count=1); return s
for ch in d.element.body.iterchildren():
    tag=ch.tag.split('}')[1]
    if tag=='p':
        p=Paragraph(ch,d); t=p.text.strip(); st=p.style.name
        if not t: continue
        if not started:
            if st.startswith('Heading') and t=='What You Will Learn in This Study': started=True
            else: continue
        if t.startswith('ZAY-MARIE PRESS DIGITAL STUDIES') or t.startswith('Copyright ©'): continue
        s=segs(p)
        if st.startswith('Heading'): sec=t; els.append(P('h1',[[t,False,False]])); continue
        if t.startswith('•'): els.append(P('bul',strip(s,r'^•\s*'))); continue
        if sec in('Study Outline','Scripture-Tracing Exercise','Review and Discussion Questions') and re.match(r'^\d+\.\s',t): els.append(P('num',strip(s,r'^\d+\.\s*'))); continue
        els.append(P('body',s))
    elif tag=='tbl' and started:
        tb=Table(ch,d)
        if len(tb.rows)==1 and len(tb.columns)==1:
            ps=[q for q in tb.cell(0,0).paragraphs if q.text.strip()]
            els.append(P('ctitle',[[ps[0].text.strip(),True,False]]))
            for q in ps[1:]: els.append(P('cbody',segs(q)))
        else:
            grid=[int(g.get(W)) for g in tb._tbl.tblGrid]; xs=[];x=0
            for g in grid: xs.append([x,x+g]); x+=g
            xs=[[round(a*640/x),round(b*640/x)] for a,b in xs]
            if tb.cell(0,0).text.strip()=='Verse': xs=[(0,135),(135,435),(435,640)]
            els.append(dict(type='tbl',rows=[[c.text.strip() for c in r.cells] for r in tb.rows],xs=xs))
for e in els:
    if e['type']=='para':
        for s in e['segs']: s[0]=s[0].replace("'","’")
    else: e['rows']=[[c.replace("'","’") for c in r] for r in e['rows']]
json.dump(els,open('els50.json','w'),ensure_ascii=False)
for i,e in enumerate(els):
    if e['type']=='tbl': print(i,'TBL',len(e['rows']),e['xs']); continue
    if e['kind'] in('h1','ctitle'): print(i,e['kind'],''.join(s[0] for s in e['segs'])[:60])
